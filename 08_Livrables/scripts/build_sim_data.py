"""Les valeurs dont le simulateur a besoin, lues dans le classeur recalculé.

Le simulateur web recalcule la trajectoire avec la même définition que le classeur
(model-def.json) ; il lui faut les valeurs des hypothèses (bornes, base, valeur active),
les multiples publiés et sectoriels, les sources des constantes, et — pour le test de
parité — ce que LibreOffice a calculé dans chaque colonne de Trajectoire et dans
Distribution. Rien n'est calculé ici : tout est lu.

Usage : python build_sim_data.py [classeur.xlsx] [model-def.json] [register.csv] [sortie.json]
"""
import csv
import json
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

HERE = Path(__file__).parent
BOOK = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / "modele-greenup-2027.xlsx"
DEF = Path(sys.argv[2]) if len(sys.argv) > 2 else BOOK.with_name("model-def.json")
REG = Path(sys.argv[3]) if len(sys.argv) > 3 else HERE / "register.csv"
OUT = Path(sys.argv[4]) if len(sys.argv) > 4 else BOOK.with_name("simulateur-data.json")

model = json.loads(DEF.read_text(encoding="utf-8"))
wb = load_workbook(BOOK, data_only=True)
register = {r["id"]: r for r in csv.DictReader(REG.open(encoding="utf-8"))}


def num(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


# ---------------------------------------------------------------- hypothèses
hyp = wb["Hypothèses"]
inputs = {}
for code, info in model["inputs"].items():
    r = info["row"]
    inputs[code] = {
        "label": info["label"], "unit": info["unit"],
        "base": num(hyp[f"C{r}"].value), "low": num(hyp[f"D{r}"].value), "high": num(hyp[f"E{r}"].value),
        "active": num(hyp[f"J{r}"].value), "why": hyp[f"G{r}"].value or "", "refs": hyp[f"H{r}"].value or "",
    }
    if inputs[code]["base"] is None:
        raise SystemExit(f"hypothèse {code} sans valeur de base : recalculer le classeur d'abord")

# ---------------------------------------------------------------- Trajectoire : entrées et sorties de chaque colonne
traj = wb["Trajectoire"]
rows_in = model["trajectoire"]["input_rows"]
rows_out = model["trajectoire"]["calc_rows"]
columns = []
for col in model["trajectoire"]["columns"]:
    c = col["col"]
    columns.append({
        "name": col["name"], "code": col["code"], "side": col["side"],
        "inputs": {code: num(traj[f"{c}{row}"].value) for code, row in rows_in.items()},
        "outputs": {key: traj[f"{c}{row}"].value for key, row in rows_out.items() if key in {x["key"] for x in model["calc"]}},
    })

# ---------------------------------------------------------------- Distribution (pour la parité du Monte-Carlo)
dist = wb["Distribution"]
distribution, probabilities = {}, {}
for r in range(1, dist.max_row + 1):
    key = dist[f"A{r}"].value
    if key in {o["key"] for o in model["outputs"]}:
        distribution[key] = {name: num(dist[f"{c}{r}"].value) for name, c in
                             (("mean", "C"), ("p5", "D"), ("p10", "E"), ("p50", "F"), ("p90", "G"), ("p95", "H"))}
    label = dist[f"B{r}"].value
    if isinstance(label, str) and num(dist[f"C{r}"].value) is not None and dist[f"A{r}"].value is None and r > 15:
        probabilities[label] = num(dist[f"C{r}"].value)
sim = wb["Simulation"]
factor_weight = num(sim["C11"].value)

# ---------------------------------------------------------------- multiples : ceux de Veolia et ceux du secteur (Cibles §D, §E)
cib = wb["Cibles"]
multiples = []
for r in range(1, cib.max_row + 1):
    label = cib[f"B{r}"].value
    if not isinstance(label, str):
        continue
    if label.startswith(("Tuck-ins", "Clean Earth,", "WTS")) and num(cib[f"C{r}"].value):
        multiples.append({"label": label, "value": num(cib[f"C{r}"].value), "basis": cib[f"D{r}"].value or "",
                          "ref": cib[f"E{r}"].value or "", "group": "Veolia"})
    if "→" in label and num(cib[f"E{r}"].value):
        multiples.append({"label": label, "value": num(cib[f"E{r}"].value), "basis": "VE / EBITDA avant synergies",
                          "ref": cib[f"H{r}"].value or "", "group": "secteur"})
        if num(cib[f"F{r}"].value):
            multiples.append({"label": label + ", après synergies", "value": num(cib[f"F{r}"].value),
                              "basis": "multiple publié après synergies", "ref": cib[f"H{r}"].value or "", "group": "secteur"})

# ---------------------------------------------------------------- levier publié 2024, 2025 (éventail de Distribution)
history = {}
for r in range(1, dist.max_row + 1):
    if dist[f"B{r}"].value in ("2024", "2025") and num(dist[f"C{r}"].value) is not None:
        history[dist[f"B{r}"].value] = num(dist[f"C{r}"].value)

# ---------------------------------------------------------------- les sources des constantes et des hypothèses
def source(fid):
    row = register.get(fid)
    if not row:
        return None
    return {"id": fid, "label": row["label"], "value": row["value"], "unit": row["unit"], "period": row["period"],
            "file": row["file"], "page": row["page"], "url": row["source_url"], "checked_by": row["checked_by"]}


consts = {name: {**info, "source": source(info["id"])} for name, info in model["consts"].items()}
for code, info in inputs.items():
    ids = [x.strip() for x in str(info["refs"]).replace(";", ",").split(",") if x.strip().startswith("F")]
    info["sources"] = [s for s in (source(i.split()[0]) for i in ids) if s]

data = {
    "generated": datetime.now(ZoneInfo("Europe/Paris")).strftime("%d/%m/%Y %H:%M"),
    "model": BOOK.name,
    "inputs": inputs, "consts": consts, "columns": columns,
    "distribution": distribution, "probabilities": probabilities, "factor_weight": factor_weight,
    "multiples": multiples, "history": history,
}
OUT.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
print("écrit", OUT, f"— {len(inputs)} hypothèses, {len(columns)} scénarios, {len(multiples)} multiples")
